# -*- coding: utf-8 -*-
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

OUT = "/home/claude/sheet/2026_렌더팜_예약시트.xlsx"

NAV   = "1F2A44"
HEAD  = PatternFill("solid", fgColor="E8ECF3")
BLOCK = PatternFill("solid", fgColor="D9D9D9")   # 수업시간 = 사용 불가
LANE  = PatternFill("solid", fgColor="F5F7FA")
OPEN_ = PatternFill("solid", fgColor="FFFFFF")
thin  = Side(style="thin", color="B9C0CC")
BOX   = Border(left=thin, right=thin, top=thin, bottom=thin)
C     = Alignment(horizontal="center", vertical="center")
L     = Alignment(horizontal="left",   vertical="center")

BLOCKS = ["11:00-15:00","15:00-19:00","19:00-23:00","23:00-03:00","03:00-07:00","07:00-11:00"]
LANES  = ["pool_a","pool_b","pool_c","pool_d"]
# 평일(월~금)은 수업 때문에 앞 두 블럭 사용 불가
CLOSED_BLOCKS = {0, 1}
WEEKDAY_COLS  = [3,4,5,6,7]           # C~G = 월~금
WEEKEND_COLS  = [8,9]                 # H,I = 토,일
ALL_COLS      = WEEKDAY_COLS + WEEKEND_COLS

MAX_TEAMS = 40
START = dt.date(2026,10,26)           # 월요일
WEEKS = 10

wb = Workbook()

# ───────────────────────────────── 명단
ms = wb.active
ms.title = "명단"
cols = [("학번",14),("팀 / 프로젝트",26),("분량(분)",10),
        ("권장 라이선스",13),("라이선스 (확정)",14),("주 최대 타임블럭",15)]
for i,(t,w) in enumerate(cols, start=1):
    c = ms.cell(row=1, column=i, value=t)
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=NAV); c.alignment = C; c.border = BOX
    ms.column_dimensions[get_column_letter(i)].width = w
for r in range(2, 2+MAX_TEAMS):
    ms.cell(row=r, column=4, value=f'=IF(C{r}="","",IF(C{r}<=3,1,IF(C{r}<=6,2,3)))')
    ms.cell(row=r, column=6, value=f'=IF(E{r}="","",E{r}*8)')
    for i in range(1,7):
        cc = ms.cell(row=r, column=i); cc.border = BOX; cc.alignment = C if i!=2 else L
ms.cell(row=MAX_TEAMS+4, column=1,
        value="※ 라이선스 1개 = 주 8 타임블럭(32시간). 권장값은 분량 기준(3분 이하 1 / 3~6분 2 / 6분 초과 3)이며, 확정은 졸준위가 입력합니다.")
ms.cell(row=MAX_TEAMS+4, column=1).font = Font(size=9, color="666666")
ms.freeze_panes = "A2"

wb.defined_names.add(DefinedName("STUDENTS", attr_text=f"명단!$A$2:$A${1+MAX_TEAMS}"))

# ───────────────────────────────── 주차 시트
def week_sheet(idx, monday):
    name = f"{monday.strftime('%m.%d')}~{(monday+dt.timedelta(days=6)).strftime('%m.%d')}"
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 9
    for col in ALL_COLS:
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["K"].width = 12

    ws["K1"] = monday
    ws["K1"].number_format = "yyyy-mm-dd"
    ws["K1"].font = Font(bold=True, color="C00000")
    ws["J1"] = "시작일(월) →"
    ws["J1"].font = Font(size=9, color="666666"); ws["J1"].alignment = Alignment(horizontal="right")
    ws["J2"] = "이 칸만 바꾸면 요일 날짜가 전부 따라옵니다"
    ws["J2"].font = Font(size=9, color="666666")

    a1 = ws.cell(row=1, column=1, value="시간대")
    b1 = ws.cell(row=1, column=2, value="레인")
    for c in (a1,b1):
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAV); c.alignment = C; c.border = BOX
    for n,col in enumerate(ALL_COLS):
        c = ws.cell(row=1, column=col, value=("=$K$1" if n==0 else f"={get_column_letter(col-1)}1+1"))
        c.number_format = 'm"/"d" ("aaa")"'
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAV); c.alignment = C; c.border = BOX

    dv = DataValidation(type="list", formula1="=STUDENTS", allow_blank=True, showDropDown=False)
    dv.error = "명단 시트에 등록된 학번만 예약할 수 있습니다."
    dv.errorTitle = "등록되지 않은 학번"
    ws.add_data_validation(dv)

    row = 2
    for bi, label in enumerate(BLOCKS):
        top = row
        ws.merge_cells(start_row=top, start_column=1, end_row=top+3, end_column=1)
        tc = ws.cell(row=top, column=1, value=label)
        tc.font = Font(bold=True, size=10); tc.alignment = C
        tc.fill = PatternFill("solid", fgColor="EEF1F6")
        for rr in range(top, top+4):
            ws.cell(row=rr, column=1).border = BOX
            lc = ws.cell(row=rr, column=2, value=LANES[rr-top])
            lc.font = Font(size=9, bold=True, color="44506B")
            lc.alignment = C; lc.border = BOX; lc.fill = LANE
            ws.row_dimensions[rr].height = 20
            for col in ALL_COLS:
                cell = ws.cell(row=rr, column=col)
                cell.border = BOX; cell.alignment = C; cell.font = Font(size=10)
                closed = (bi in CLOSED_BLOCKS and col in WEEKDAY_COLS)
                if closed:
                    cell.fill = BLOCK
                else:
                    cell.fill = OPEN_
                    dv.add(cell)
        row += 4

    note = ws.cell(row=row+1, column=1,
                   value="■ 회색 칸은 수업 시간이라 사용할 수 없습니다 (평일 11:00~19:00). 예약은 본인 학번을 드롭다운에서 고르세요.")
    note.font = Font(size=9, color="666666")

    # ── 남은 타임블럭
    h = row + 3
    heads = ["학번","주 최대","사용","남은"]
    for i,t in enumerate(heads, start=1):
        c = ws.cell(row=h, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAV); c.alignment = C; c.border = BOX
    first, last = 2, 25
    for k in range(MAX_TEAMS):
        r = h+1+k
        mr = 2+k
        ws.cell(row=r, column=1, value=f'=IF(명단!A{mr}="","",명단!A{mr})')
        ws.cell(row=r, column=2, value=f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},명단!$A:$F,6,0),""))')
        ws.cell(row=r, column=3, value=f'=IF($A{r}="","",COUNTIF($C${first}:$I${last},$A{r}))')
        ws.cell(row=r, column=4, value=f'=IF($A{r}="","",IFERROR($B{r}-$C{r},""))')
        for i in range(1,5):
            cc = ws.cell(row=r, column=i); cc.border = BOX; cc.alignment = C
    ws.cell(row=h+MAX_TEAMS+2, column=1,
            value="※ 이 표는 자동입니다. 학번과 라이선스는 [명단] 시트에서만 관리하세요.").font = Font(size=9, color="666666")

    ws.freeze_panes = "C2"
    return ws

for i in range(WEEKS):
    week_sheet(i, START + dt.timedelta(weeks=i))

# ───────────────────────────────── 규칙
rs = wb.create_sheet("사용 규칙")
rs.column_dimensions["A"].width = 30
rs.column_dimensions["B"].width = 62
rows = [
 ("타임블럭", "4시간 단위. 평일 19:00~익일 11:00 (4블럭) · 주말 24시간 (6블럭)"),
 ("레인", "실습실 PC를 4그룹으로 나눈 단위. 한 시간대에 4팀이 동시에 렌더합니다"),
 ("라이선스", "소프트웨어 라이선스가 아니라 주당 예약 상한입니다. 1개 = 주 8블럭(32시간)"),
 ("주간 공급량", "평일 5일×4블럭×4레인(80) + 주말 2일×6블럭×4레인(48) = 128블럭"),
 ("예약 방법", "본인 주차 시트에서 시간대+레인 칸을 고르고 드롭다운에서 본인 학번 선택"),
 ("제출할 때", "Deadline 제출 창의 Pool 에 예약한 레인(pool_a~pool_d)을 그대로 지정"),
 ("", ""),
 ("❌ 위반 행위", ""),
 ("타 학번으로 예약", "본인 팀이 아닌 학번으로 블럭 선점"),
 ("타인의 예약 수정/가로채기", "이미 예약된 칸을 자신의 학번으로 변경하거나 삭제"),
 ("고의적 리소스 소모", "불필요한 예약/취소 반복, 테스트 목적이 아닌 무분별 렌더"),
 ("", ""),
 ("🛑 처리 수위", ""),
 ("1회", "경고 및 기록 공유"),
 ("2회", "해당 주 라이선스 -2블럭"),
 ("3회", "3일간 렌더팜 사용 정지 + 교수/지도"),
 ("4회 이상", "팀 렌더팜 사용권 박탈"),
 ("", ""),
 ("🤝 부탁", "렌더팜은 공동 자산이며, 모두의 졸업 프로젝트 완성을 위해 운영됩니다."),
 ("", "문제가 생기거나 양보가 필요할 때는 항상 대화로 해결해주세요."),
]
for r,(a,b) in enumerate(rows, start=1):
    ca = rs.cell(row=r, column=1, value=a); ca.alignment = L
    cb = rs.cell(row=r, column=2, value=b); cb.alignment = L
    if a.startswith(("❌","🛑","🤝")) or (a and not b):
        ca.font = Font(bold=True, size=11)
    elif a:
        ca.font = Font(bold=True, size=10)
    rs.row_dimensions[r].height = 19

wb.save(OUT)
print("saved:", OUT)
print("sheets:", wb.sheetnames)
