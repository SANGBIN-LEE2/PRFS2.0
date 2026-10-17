import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

NAVY, ACCENT, HEAD = "1F3251", "C0392B", "2F4F76"
SOFT, BAND, NIGHT, BLOCKED = "EAEFF5", "F7F9FC", "FFF4E0", "DDE3EA"
GRP = {"A": "E8F0FB", "B": "EAF6EC", "C": "FDF0E6", "D": "F2ECF8"}

thin = Side(style="thin", color="C8D2DD")
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def head(ws, row, values, start=1, fill=HEAD, size=10):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start + i, value=v)
        c.font = Font(bold=True, color="FFFFFF", size=size)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box


def title(ws, row, text, sub=None):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(bold=True, size=16, color=NAVY)
    ws.row_dimensions[row].height = 26
    if sub:
        s = ws.cell(row=row + 1, column=1, value=sub)
        s.font = Font(size=10, color="5A6B7D")


wb = Workbook()

# ══════════════════════════════ 01 신청서
ws = wb.active
ws.title = "01_신청서"
ws.sheet_view.showGridLines = False
title(ws, 1, "M211 렌더팜 사용 신청서",
      "팀별로 한 줄씩 작성해주세요.  분량은 최종 상영 길이(분) 기준입니다.  회색 칸은 자동 계산되니 건드리지 마세요.")

cols = ["No.", "팀 프로젝트명", "팀장 학번", "팀장 이메일", "프로젝트 분량(분)",
        "라이선스", "주 타임블럭", "주 사용시간", "신청일", "비고"]
head(ws, 4, cols)
ws.row_dimensions[4].height = 32
for i, w in enumerate([6, 28, 13, 28, 15, 10, 12, 12, 12, 24], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

FIRST, LAST = 5, 34
for r in range(FIRST, LAST + 1):
    ws.cell(row=r, column=1, value=r - FIRST + 1)
    ws.cell(row=r, column=6, value=f'=IF(E{r}="","",IF(E{r}<=3,1,IF(E{r}<=6,2,3)))')
    ws.cell(row=r, column=7, value=f'=IF(F{r}="","",F{r}*8)')
    ws.cell(row=r, column=8, value=f'=IF(G{r}="","",G{r}*4)')
    for c in range(1, 11):
        cell = ws.cell(row=r, column=c)
        cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if c in (6, 7, 8):
            cell.fill = PatternFill("solid", fgColor=SOFT)
            cell.font = Font(bold=True, color=NAVY)
        elif r % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=BAND)
    for c in (2, 4, 10):
        ws.cell(row=r, column=c).alignment = Alignment(horizontal="left", vertical="center")

n = LAST + 2
ws.cell(row=n, column=2, value="합계").font = Font(bold=True, color=NAVY)
ws.cell(row=n, column=6, value=f"=SUM(F{FIRST}:F{LAST})").font = Font(bold=True)
ws.cell(row=n, column=7, value=f"=SUM(G{FIRST}:G{LAST})").font = Font(bold=True, color=ACCENT)
ws.cell(row=n, column=8, value=f"=SUM(H{FIRST}:H{LAST})").font = Font(bold=True)

chk = ws.cell(row=n + 1, column=2,
              value=f'=IF(G{n}="","",IF(G{n}<=128,"수용 가능  —  주간 공급 128블럭 중 "&G{n}&"블럭 신청",'
                    f'"초과  —  주간 공급 128블럭보다 "&(G{n}-128)&"블럭 많습니다. 산정 기준 조정 필요"))')
chk.font = Font(bold=True, size=11, color=NAVY)

ws.conditional_formatting.add(f"G{n}", CellIsRule(
    operator="greaterThan", formula=["128"],
    fill=PatternFill("solid", fgColor="FCE4E4"), font=Font(bold=True, color="9B1C1C")))

info = [
    ("", ""),
    ("분량 → 라이선스 산정 기준", ""),
    ("3분 이하", "라이선스 1  ·  주 8블럭  ·  32시간"),
    ("3분 초과 ~ 6분 이하", "라이선스 2  ·  주 16블럭  ·  64시간"),
    ("6분 초과", "라이선스 3  ·  주 24블럭  ·  96시간"),
    ("", ""),
    ("주간 공급량", "평일 5일×4블럭×4그룹 = 80  +  주말 2일×6블럭×4그룹 = 48  →  총 128블럭"),
]
base = n + 3
for i, (a, b) in enumerate(info):
    ca, cb = ws.cell(row=base + i, column=2, value=a), ws.cell(row=base + i, column=3, value=b)
    ca.font = Font(bold=True, size=11, color=NAVY) if i in (1, 6) else Font(size=10)
    cb.font = Font(size=10, color="5A6B7D")

dv = DataValidation(type="decimal", operator="between", formula1=0.1, formula2=60,
                    showErrorMessage=True, errorTitle="분량 확인",
                    error="프로젝트 분량을 분 단위 숫자로 입력하세요. (예: 4.5)")
ws.add_data_validation(dv)
dv.add(f"E{FIRST}:E{LAST}")
ws.freeze_panes = "A5"

# ══════════════════════════════ 02 예약표
ws2 = wb.create_sheet("02_예약표")
ws2.sheet_view.showGridLines = False
title(ws2, 1, "예약표",
      "빈 칸에 본인 학번을 입력하세요.  회색 칸은 수업 시간이라 사용할 수 없습니다.  "
      "제출할 때 반드시 예약한 그룹을 지정하세요.")

BLOCKS = ["19:00-23:00", "23:00-03:00", "03:00-07:00", "07:00-11:00",
          "11:00-15:00", "15:00-19:00"]
WEEKEND_ONLY = {"11:00-15:00", "15:00-19:00"}
NIGHTY = {"23:00-03:00", "03:00-07:00"}
DAYS = ["월", "화", "수", "목", "금", "토", "일"]
GROUPS = ["A", "B", "C", "D"]

ws2.column_dimensions["A"].width = 7
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 7
for i in range(4, 11):
    ws2.column_dimensions[get_column_letter(i)].width = 13

start_monday = dt.date(2026, 10, 19)
row, week_ranges = 4, []
for wk in range(4):
    mon = start_monday + dt.timedelta(days=7 * wk)
    sun = mon + dt.timedelta(days=6)
    t = ws2.cell(row=row, column=1,
                 value=f"{wk+1}주차   {mon.strftime('%m/%d')} ~ {sun.strftime('%m/%d')}")
    t.font = Font(bold=True, size=12, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    ws2.row_dimensions[row].height = 22
    row += 1

    head(ws2, row, ["주차", "타임블럭", "그룹"] +
         [f"{d}\n{(mon + dt.timedelta(days=i)).strftime('%m/%d')}" for i, d in enumerate(DAYS)])
    ws2.row_dimensions[row].height = 30
    row += 1

    first = row
    for b in BLOCKS:
        for g in GROUPS:
            ws2.cell(row=row, column=1, value=f"{wk+1}주")
            bc = ws2.cell(row=row, column=2, value=b if g == "A" else "")
            bc.font = Font(bold=True, size=10, color=NAVY)
            gc = ws2.cell(row=row, column=3, value=g)
            gc.font = Font(bold=True, size=10, color=NAVY)
            gc.fill = PatternFill("solid", fgColor=GRP[g])
            for c in range(1, 11):
                cell = ws2.cell(row=row, column=c)
                cell.border = box
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if c == 1 or c == 2:
                    cell.fill = PatternFill("solid", fgColor=SOFT)
                elif c == 3:
                    pass
                elif b in WEEKEND_ONLY and c <= 8:          # 월~금 = col 4..8
                    cell.fill = PatternFill("solid", fgColor=BLOCKED)
                    if g == "A" and c == 4:
                        cell.value = "수업 시간"
                        cell.font = Font(size=9, color="7A8696", italic=True)
                elif b in NIGHTY:
                    cell.fill = PatternFill("solid", fgColor=NIGHT)
            ws2.row_dimensions[row].height = 19
            row += 1
    week_ranges.append((first, row - 1))
    row += 1

ws2.cell(row=row, column=1,
         value="노란 칸은 심야 블록입니다. 가장 한가하니 적극 활용하세요 — 렌더는 지켜보지 않아도 돌아갑니다."
         ).font = Font(size=10, color="8A6D1F")
ws2.cell(row=row + 1, column=1,
         value="그룹 A·B·C·D 는 서로 다른 PC 묶음입니다. 같은 시간대에 네 팀이 동시에 렌더할 수 있습니다."
         ).font = Font(size=10, color="5A6B7D")
ws2.freeze_panes = "D5"

# ══════════════════════════════ 03 사용현황
ws3 = wb.create_sheet("03_사용현황")
ws3.sheet_view.showGridLines = False
title(ws3, 1, "사용 현황",
      "신청서와 예약표에서 자동으로 계산됩니다.  남은 블럭이 빨간색이면 상한을 초과한 것입니다.")

head(ws3, 4, ["팀 프로젝트명", "학번", "주 최대"] +
     [f"{i+1}주차 사용" for i in range(4)] + [f"{i+1}주차 남음" for i in range(4)])
ws3.row_dimensions[4].height = 30
for i, w in enumerate([26, 13, 10, 11, 11, 11, 11, 11, 11, 11, 11], start=1):
    ws3.column_dimensions[get_column_letter(i)].width = w

for i in range(30):
    r, src = 5 + i, FIRST + i
    ws3.cell(row=r, column=1, value=f"=IF('01_신청서'!B{src}=\"\",\"\",'01_신청서'!B{src})")
    ws3.cell(row=r, column=2, value=f"=IF('01_신청서'!C{src}=\"\",\"\",'01_신청서'!C{src})")
    ws3.cell(row=r, column=3, value=f"=IF('01_신청서'!G{src}=\"\",\"\",'01_신청서'!G{src})")
    for w, (a, b) in enumerate(week_ranges):
        rng = f"'02_예약표'!$D${a}:$J${b}"
        ws3.cell(row=r, column=4 + w, value=f'=IF($B{r}="","",COUNTIF({rng},$B{r}))')
        ws3.cell(row=r, column=8 + w,
                 value=f'=IF($B{r}="","",$C{r}-{get_column_letter(4+w)}{r})')
    for c in range(1, 12):
        cell = ws3.cell(row=r, column=c)
        cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if c >= 8:
            cell.fill = PatternFill("solid", fgColor=SOFT)
    ws3.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")

tot = 36
ws3.cell(row=tot, column=1, value="주차별 전체 예약 블럭").font = Font(bold=True, color=NAVY)
for w, (a, b) in enumerate(week_ranges):
    rng = f"'02_예약표'!$D${a}:$J${b}"
    c = ws3.cell(row=tot, column=4 + w, value=f'=COUNTA({rng})-COUNTIF({rng},"수업 시간")')
    c.font = Font(bold=True, color=ACCENT)
    c.alignment = Alignment(horizontal="center")
ws3.cell(row=tot + 1, column=1, value="주간 공급량 128블럭").font = Font(size=10, color="5A6B7D")

ws3.conditional_formatting.add("H5:K34", CellIsRule(
    operator="lessThan", formula=["0"],
    fill=PatternFill("solid", fgColor="FCE4E4"), font=Font(color="9B1C1C", bold=True)))
ws3.conditional_formatting.add("H5:K34", CellIsRule(
    operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="FFF6E0")))
ws3.freeze_panes = "A5"

# ══════════════════════════════ 04 안내
ws4 = wb.create_sheet("04_안내")
ws4.sheet_view.showGridLines = False
ws4.column_dimensions["A"].width = 3
ws4.column_dimensions["B"].width = 30
ws4.column_dimensions["C"].width = 76

lines = [
    ("H", "렌더팜 사용 안내", ""),
    ("P", "", "자세한 사용법은 노션 가이드를 참고하세요. 이 시트는 신청과 예약만 다룹니다."),
    ("", "", ""),
    ("H2", "1.  신청", ""),
    ("K", "제출 항목", "팀 프로젝트명 / 팀장 학번 / 팀장 이메일 / 프로젝트 분량(분)"),
    ("K", "분량 기준", "최종 상영 길이. 렌더해야 할 전체 컷의 길이를 적어주세요."),
    ("K", "접수", "01_신청서 탭에 팀별로 한 줄씩 작성"),
    ("", "", ""),
    ("H2", "2.  타임블럭과 그룹", ""),
    ("K", "타임블럭 1개", "4시간"),
    ("K", "평일 (월~금)", "19:00 ~ 11:00 만 사용 가능  →  하루 4블럭  (낮 11~19시는 수업)"),
    ("K", "주말 (토·일)", "24시간  →  하루 6블럭"),
    ("K", "PC 그룹", "A · B · C · D 네 묶음. 한 시간대에 네 팀이 동시에 렌더"),
    ("K", "주간 공급량", "평일 5×4×4 = 80  +  주말 2×6×4 = 48  →  총 128블럭"),
    ("", "", ""),
    ("H2", "3.  라이선스 배정", ""),
    ("K", "라이선스 1개", "주 8블럭 = 32시간"),
    ("K", "3분 이하", "라이선스 1  ·  주 8블럭"),
    ("K", "3분 초과 ~ 6분 이하", "라이선스 2  ·  주 16블럭"),
    ("K", "6분 초과", "라이선스 3  ·  주 24블럭"),
    ("P", "", "라이선스 합계가 16을 넘으면(128블럭 초과) 산정 기준을 조정해야 합니다."),
    ("P", "", "라이선스는 상한선이며 할당량이 아닙니다. 블록은 선착순입니다."),
    ("", "", ""),
    ("H2", "4.  예약", ""),
    ("K", "방법", "02_예약표에서 시간 + 그룹을 고르고 빈 칸에 본인 학번 입력"),
    ("K", "제출할 때", "Deadline 제출 창의 Pool / Group 을 예약한 그룹과 똑같이 지정"),
    ("K", "미등록 학번", "자동 삭제됩니다"),
    ("K", "안 쓰는 블록", "미리 비워주세요. 그래야 다른 팀이 씁니다"),
    ("K", "심야 블록", "23-03 / 03-07 이 가장 한가합니다. 저녁에 제출하고 자면 됩니다"),
    ("K", "PC 로그인", "예약한 시간에만 Worker PC를 사용하세요"),
    ("", "", ""),
    ("H2", "5.  위반 시", ""),
    ("K", "위반 유형", "타 학번으로 예약 / 타인 예약 수정·가로채기 / 예약 없이 제출 / 블록 선점 후 방치"),
    ("K", "1회", "경고 및 기록 공유"),
    ("K", "2회", "해당 주 라이선스 −2블럭"),
    ("K", "3회", "3일간 렌더팜 사용 정지"),
    ("K", "4회 이상", "팀 렌더팜 사용권 박탈 (졸준위·지도교수 협의)"),
    ("", "", ""),
    ("H2", "6.  문의", ""),
    ("K", "예약 · 타임블럭 배정", "졸업준비위원회"),
    ("K", "렌더 실패 · 장비 문제", "이상빈 (C393169)"),
]
r = 2
for kind, a, b in lines:
    ca, cb = ws4.cell(row=r, column=2, value=a), ws4.cell(row=r, column=3, value=b)
    if kind == "H":
        ca.font = Font(bold=True, size=17, color=NAVY)
        ws4.row_dimensions[r].height = 28
    elif kind == "H2":
        ca.font = Font(bold=True, size=12, color="FFFFFF")
        ca.fill = PatternFill("solid", fgColor=HEAD)
        cb.fill = PatternFill("solid", fgColor=HEAD)
        ca.alignment = Alignment(vertical="center", indent=1)
        ws4.row_dimensions[r].height = 22
    elif kind == "K":
        ca.font = Font(bold=True, size=10, color=NAVY)
        cb.font = Font(size=10)
        ca.alignment = Alignment(vertical="center", indent=1)
        cb.alignment = Alignment(vertical="center", wrap_text=True)
        ws4.row_dimensions[r].height = 19
    elif kind == "P":
        cb.font = Font(size=10, italic=True, color=ACCENT)
        cb.alignment = Alignment(vertical="center", wrap_text=True)
    r += 1

out = "/home/claude/M211_렌더팜_신청_예약시트.xlsx"
wb.save(out)
print("saved", out)
